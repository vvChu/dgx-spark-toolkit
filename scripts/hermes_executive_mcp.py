#!/usr/bin/env python3
"""Hermes Executive MCP Server.

Provides safe, sandboxed, high-level tools for Hermes Agent on Telegram:
1. read_cached_document: Convert and extract .doc, .docx, .pdf, .xlsx from cache/documents/
2. search_legal_corpus: Query Vietnamese Legal RAG (Milvus legal_docs_v11)
3. validate_legal_citation: Inspect citation validity and genealogy via Neo4j
4. get_system_health: Query live GPU, memory, container health probes via ChatOps :8095
5. propose_system_operation: Submit operational commands for Leader 1-tap Telegram approval
"""

from __future__ import annotations

import json
import math
import os
from pathlib import Path
import subprocess
import sys
from typing import Any, Dict, Optional
from urllib.parse import quote

try:
    from mcp.server.mcpserver import MCPServer
except ImportError:
    MCPServer = None

try:
    import httpx
except ImportError:
    httpx = None


DOCS_CACHE_DIR = Path(os.path.expanduser("~/.hermes/cache/documents")).resolve()
RAG_SERVICE_URL = os.environ.get("RAG_SERVICE_URL", "http://localhost:8005")
CHATOPS_BASE_URL = os.environ.get("CHATOPS_BASE_URL", "http://127.0.0.1:8095")
CHATOPS_SECRET = os.environ.get("CHATOPS_INTERNAL_SECRET", "").strip()

ALLOWED_MUTATING_COMMANDS = frozenset([
    "system.container.restart",
    "system.openwebui.upgrade",
    "antigravity.account.reenable",
    "system.deps.check",
    "system.deps.upgrade",
    "ccba.skill.boost",
])


def _validate_safe_path(p: Path) -> Optional[str]:
    """Xác thực đường dẫn an toàn: nằm trong cache dir, không phải symlink, không phải hardlink."""
    # 1. Kiểm tra symlink trước khi resolve (lstat)
    if p.is_symlink() or os.path.islink(p):
        return "Security blocked: Không cho phép truy cập symbolic link."

    # 2. Kiểm tra hardlink trước khi resolve
    try:
        if p.exists() and p.lstat().st_nlink > 1:
            return "Security blocked: Không cho phép truy cập hardlink (st_nlink > 1)."
    except OSError as e:
        return f"Security blocked: Lỗi đọc metadata file: {e}"

    # 3. Resolve và kiểm tra nằm trong DOCS_CACHE_DIR
    try:
        resolved = p.resolve()
        base_dir = DOCS_CACHE_DIR.resolve()
        if not resolved.is_relative_to(base_dir):
            return "Path traversal blocked: Tệp phải nằm trong ~/.hermes/cache/documents/"
    except Exception:
        return "Path traversal blocked: Không thể xác định đường dẫn hợp lệ."

    # 4. Kiểm tra lại sau resolve (chặn hardlink đích)
    try:
        if resolved.exists() and resolved.stat().st_nlink > 1:
            return "Security blocked: Không cho phép truy cập hardlink (st_nlink > 1)."
    except OSError:
        pass

    return None


def _extract_doc_or_docx(path: Path) -> str:
    """Extracts text and tables from .doc or .docx into Markdown."""
    target_docx = path
    if path.suffix.lower() == ".doc":
        expected_docx = path.with_suffix(".docx")

        # Kiểm tra trước nếu expected_docx đã tồn tại trong cache
        if expected_docx.exists() or expected_docx.is_symlink() or os.path.islink(expected_docx):
            err = _validate_safe_path(expected_docx)
            if err:
                return f"[{err}]"

        if not expected_docx.exists() or expected_docx.stat().st_mtime < path.stat().st_mtime:
            # Convert via LibreOffice headless
            env = os.environ.copy()
            env["SAL_USE_VCLPLUGIN"] = "svp"
            res = subprocess.run(
                [
                    "soffice",
                    "--headless",
                    "--convert-to",
                    "docx",
                    str(path),
                    "--outdir",
                    str(path.parent),
                ],
                env=env,
                capture_output=True,
                text=True,
                timeout=60,
            )
            if res.returncode != 0:
                return f"[Lỗi chuyển đổi LibreOffice]: {res.stderr or res.stdout}"

        # Kiểm tra nghiêm ngặt file docx sinh ra sau khi convert
        err = _validate_safe_path(expected_docx)
        if err:
            return f"[{err}]"

        target_docx = expected_docx

    # Xác thực target_docx an toàn trước khi nạp vào python-docx
    err = _validate_safe_path(target_docx)
    if err:
        return f"[{err}]"

    import docx

    doc = docx.Document(target_docx)
    lines: list[str] = [f"# {path.name}", ""]
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            lines.append(t)

    for idx, table in enumerate(doc.tables):
        lines.append(f"\n### Bảng {idx + 1}")
        for row in table.rows:
            cells = [c.text.replace("\n", " ").strip() for c in row.cells]
            lines.append("| " + " | ".join(cells) + " |")

    return "\n".join(lines)


def _extract_pdf(path: Path) -> str:
    """Extracts text from PDF via PyMuPDF (fitz) or pypdf."""
    try:
        import fitz  # PyMuPDF

        doc = fitz.open(path)
        lines: list[str] = [f"# {path.name}", f"- Số trang: {len(doc)}", ""]
        for p_no in range(len(doc)):
            page = doc[p_no]
            t = page.get_text().strip()
            if t:
                lines.append(f"## Trang {p_no + 1}\n{t}\n")
        return "\n".join(lines)
    except Exception:
        import pypdf

        reader = pypdf.PdfReader(str(path))
        lines = [f"# {path.name}", f"- Số trang: {len(reader.pages)}", ""]
        for p_no, page in enumerate(reader.pages):
            t = page.extract_text() or ""
            if t.strip():
                lines.append(f"## Trang {p_no + 1}\n{t.strip()}\n")
        return "\n".join(lines)


def _extract_xlsx(path: Path) -> str:
    """Extracts spreadsheet rows into Markdown tables."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    lines: list[str] = [f"# {path.name}", ""]
    for sheetname in wb.sheetnames:
        sheet = wb[sheetname]
        lines.append(f"## Sheet: {sheetname}")
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                cells = [str(c).replace("\n", " ").strip() if c is not None else "" for c in row]
                lines.append("| " + " | ".join(cells) + " |")
        lines.append("")
    return "\n".join(lines)


def _paginate(text: str, page: int, page_size: int = 10000) -> str:
    """Paginates text to prevent Telegram context window explosion."""
    page_size = min(max(page_size, 2000), 12000)
    total_len = len(text)
    if total_len <= page_size and page == 1:
        return text

    total_pages = max(1, math.ceil(total_len / page_size))
    page = max(1, min(page, total_pages))
    start = (page - 1) * page_size
    end = min(start + page_size, total_len)

    sliced = text[start:end]
    footer = (
        f"\n\n---\n"
        f"📄 *[Trang {page}/{total_pages} | Ký tự {start + 1}-{end}/{total_len}]*\n"
        f"💡 *Gọi lại tool với `page={page + 1}` để đọc tiếp trang sau.*"
    )
    return sliced + footer


# --- Function implementations exposed to tool registry ---
def read_cached_document(filename: str, page: int = 1, page_size: int = 10000) -> str:
    """Đọc và trích xuất cấu trúc văn bản (.doc, .docx, .pdf, .xlsx) từ thư mục cache tài liệu.

    Args:
        filename: Tên tệp trong thư mục ~/.hermes/cache/documents/ (ví dụ: 'doc_123_Mau_so_2.doc').
        page: Số trang cần đọc (mặc định 1).
        page_size: Số ký tự tối đa trên mỗi trang (mặc định 10.000, trần 12.000).

    Returns:
        Nội dung văn bản và bảng biểu dưới dạng Markdown sạch đã được phân trang an toàn.
    """
    fn = filename.strip()
    if "/" in fn or "\\" in fn or ".." in fn:
        return json.dumps({"error": "Path traversal blocked: Tên tệp không được chứa ký tự phân cách thư mục hoặc '..'"})

    raw_path = DOCS_CACHE_DIR / fn
    err = _validate_safe_path(raw_path)
    if err:
        return json.dumps({"error": err})

    target_path = raw_path.resolve()
    if not target_path.exists() or not target_path.is_file():
        return json.dumps({"error": f"Không tìm thấy tệp '{fn}' trong cache documents."})

    ext = target_path.suffix.lower()
    try:
        if ext in (".doc", ".docx"):
            full_text = _extract_doc_or_docx(target_path)
        elif ext == ".pdf":
            full_text = _extract_pdf(target_path)
        elif ext == ".xlsx":
            full_text = _extract_xlsx(target_path)
        elif ext in (".txt", ".md", ".json", ".yaml", ".yml", ".csv"):
            full_text = target_path.read_text(encoding="utf-8", errors="replace")
        else:
            return json.dumps({"error": f"Định dạng '{ext}' chưa được hỗ trợ bóc tách tự động."})

        if full_text.startswith("[Security blocked:") or full_text.startswith("[Path traversal blocked:"):
            return json.dumps({"error": full_text.strip("[]")})

        return _paginate(full_text, page=page, page_size=page_size)
    except Exception as e:
        return json.dumps({"error": f"Lỗi bóc tách tài liệu {target_path.name}: {e}"})


def search_legal_corpus(query: str, limit: int = 5, use_reranker: bool = True) -> str:
    """Tra cứu văn bản quy phạm pháp luật Việt Nam (QCVN, TCVN, Luật Xây dựng, Nghị định) từ Milvus."""
    if httpx is None:
        return json.dumps({"error": "httpx package is missing"})
    safe_limit = min(max(int(limit), 1), 8)
    try:
        with httpx.Client(timeout=30.0) as client:
            res = client.post(
                f"{RAG_SERVICE_URL}/search",
                json={"query": query.strip(), "limit": safe_limit, "use_reranker": use_reranker},
            )
            if res.status_code == 200:
                return res.text
            return json.dumps({"error": f"RAG Service HTTP {res.status_code}", "detail": res.text})
    except Exception as e:
        return json.dumps({"error": f"Không thể kết nối RAG Service ({RAG_SERVICE_URL}): {e}"})


def validate_legal_citation(so_hieu: str) -> str:
    """Tra cứu tính hợp lệ và phả hệ sửa đổi/thay thế của văn bản từ Neo4j Legal Graph."""
    if httpx is None:
        return json.dumps({"error": "httpx package is missing"})
    clean_id = so_hieu.strip()
    try:
        encoded_id = quote(clean_id, safe="")
        with httpx.Client(timeout=10.0) as client:
            res = client.get(f"{RAG_SERVICE_URL}/graph/neighbors?id={encoded_id}")
            if res.status_code == 200:
                return res.text
            return json.dumps({"error": f"Legal Graph HTTP {res.status_code}", "detail": res.text})
    except Exception as e:
        return json.dumps({"error": f"Không thể kết nối Neo4j Graph Service: {e}"})


def get_system_health(probe_type: str = "gpu") -> str:
    """Tra cứu số liệu sức khỏe hạ tầng DGX Spark thời gian thực qua ChatOps Daemon."""
    if httpx is None:
        return json.dumps({"error": "httpx package is missing"})
    secret = CHATOPS_SECRET or os.environ.get("CHATOPS_INTERNAL_SECRET", "").strip()
    if not secret:
        return json.dumps({"error": "CHATOPS_INTERNAL_SECRET chưa được cấu hình trong môi trường."})
    clean_probe = probe_type.lower().strip()
    try:
        headers = {"X-ChatOps-Secret": secret}
        with httpx.Client(timeout=15.0) as client:
            res = client.get(f"{CHATOPS_BASE_URL}/api/v1/probe/{clean_probe}", headers=headers)
            if res.status_code == 200:
                data = res.json()
                return data.get("result", str(data))
            return json.dumps({"error": f"ChatOps Probe HTTP {res.status_code}", "detail": res.text})
    except Exception as e:
        return json.dumps({"error": f"Không thể kết nối ChatOps Daemon ({CHATOPS_BASE_URL}): {e}"})


def propose_system_operation(command_id: str, reason: str, params: Optional[Dict[str, Any]] = None) -> str:
    """Gửi đề xuất thực thi tác vụ hệ thống sang Telegram để Lãnh đạo phê duyệt 1-chạm."""
    if httpx is None:
        return json.dumps({"error": "httpx package is missing"})

    cmd = command_id.strip()
    if cmd not in ALLOWED_MUTATING_COMMANDS:
        allowed_list = ", ".join(sorted(ALLOWED_MUTATING_COMMANDS))
        return json.dumps({
            "error": f"Lệnh '{cmd}' không nằm trong danh mục đề xuất an toàn.",
            "allowed_commands": allowed_list,
        })

    secret = CHATOPS_SECRET or os.environ.get("CHATOPS_INTERNAL_SECRET", "").strip()
    if not secret:
        return json.dumps({"error": "CHATOPS_INTERNAL_SECRET chưa được cấu hình trong môi trường."})

    param_dict = params or {}
    payload = {
        "title": f"Đề xuất Vận hành: {cmd}",
        "body": (
            f"🤖 *Hermes đề xuất thực thi tác vụ hệ thống:*\n"
            f"• *Mã lệnh:* `{cmd}`\n"
            f"• *Lý do:* {reason.strip()}\n"
            f"• *Tham số:* `{json.dumps(param_dict, ensure_ascii=False)}`\n\n"
            f"_Vui lòng bấm nút bên dưới để xác nhận thực thi hoặc bỏ qua nếu không đồng ý._"
        ),
        "severity": "WARNING",
        "actions": [
            {
                "action_id": "approve_ops",
                "label": f"✅ Phê duyệt {cmd}",
                "command": cmd,
                "params": param_dict,
                "ttl_seconds": 300,
            }
        ],
    }

    try:
        headers = {"X-ChatOps-Secret": secret}
        with httpx.Client(timeout=15.0) as client:
            res = client.post(f"{CHATOPS_BASE_URL}/api/v1/notify", json=payload, headers=headers)
            if res.status_code == 200:
                data = res.json()
                msg_id = data.get("telegram_message_id", "N/A")
                return (
                    f"✅ Đã gửi Thẻ phê duyệt tác vụ `{cmd}` sang Telegram của Lãnh đạo (Telegram Msg ID: {msg_id}).\n"
                    f"Tác vụ đang ở trạng thái chờ Lãnh đạo bấm nút 'Phê duyệt' trên điện thoại."
                )
            return json.dumps({"error": f"ChatOps Dispatch HTTP {res.status_code}", "detail": res.text})
    except Exception as e:
        return json.dumps({"error": f"Không thể gửi đề xuất sang ChatOps: {e}"})


def create_executive_mcp_server() -> Any:
    if MCPServer is None:
        raise RuntimeError("Gói 'mcp' chưa được cài đặt. Vui lòng cài: pip install mcp>=2.2.0")

    server = MCPServer(name="hermes-executive-mcp")
    server.tool()(read_cached_document)
    server.tool()(search_legal_corpus)
    server.tool()(validate_legal_citation)
    server.tool()(get_system_health)
    server.tool()(propose_system_operation)
    return server


def main() -> int:
    if MCPServer is None:
        print("[ERROR] Gói 'mcp' chưa được cài đặt trong môi trường Python.", file=sys.stderr)
        return 1

    server = create_executive_mcp_server()
    server.run(transport="stdio")
    return 0


if __name__ == "__main__":
    sys.exit(main())
